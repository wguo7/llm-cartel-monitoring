"""Fill the official Apart report template (template/template.docx) with report/report.md -> report/report.docx.

The template's header block (title, author grid, "With Apart Research", abstract box) is kept and filled in place; the
guidance box and guidance text are removed and the report body is appended in the template's fonts and heading styles.
Usage: python scripts/build_docx.py   (then: soffice --headless --convert-to pdf report/report.docx)
"""
import copy
import pathlib
import re

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = pathlib.Path(__file__).resolve().parents[1]
FONT = "Old Standard TT"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
NUM_BULLET, NUM_CONTRIB, NUM_REFS = 3, 5, 1  # numbering ids defined in the template

md = (ROOT / "report/report.md").read_text().replace("&lt;", "<").replace("&gt;", ">")
d = docx.Document(ROOT / "template/template.docx")
body = d.element.body


def set_font(run, size=None, bold=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    rpr.get_or_add_rFonts().set(qn("w:cs"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(p, text, size=None, italic=False):
    """Add text with **bold** and *italic* spans."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            set_font(p.add_run(tok[2:-2]), size, bold=True, italic=italic or None)
        elif tok.startswith("*"):
            set_font(p.add_run(tok[1:-1]), size, italic=True)
        else:
            set_font(p.add_run(tok), size, italic=italic or None)


def spacing(p, before=0, after=8, line=1.15):
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line


def numbered(p, num_id):
    ppr = p._element.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl, nid = OxmlElement("w:ilvl"), OxmlElement("w:numId")
    ilvl.set(qn("w:val"), "0")
    nid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(nid)
    ppr.append(numpr)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "300")
    ppr.append(ind)


# Header block: title, author, abstract.
hdr = d.tables[0]
title = re.search(r"^# (.+)$", md, re.M).group(1)
runs = [r for p in hdr.rows[0].cells[0].paragraphs for r in p.runs if r.text.strip()]
runs[0].text = title
for r in runs[1:]:
    r.text = ""
grid = [t for t in hdr._element.iter(W + "tbl") if len(t.findall(W + "tr")) == 2 and len(t.find(W + "tr").findall(W + "tc")) == 3][0]
rows = grid.findall(W + "tr")
for k, tc in enumerate(rows[0].findall(W + "tc")):
    ts = list(tc.iter(W + "t"))
    if k == 1:
        ts[0].text, ts[1].text = "William Guo", "University of Chicago"
    else:
        for t in ts:
            t.text = ""
        for br in list(tc.iter(W + "br")):
            br.getparent().remove(br)
grid.remove(rows[1])
abstract = re.search(r"## Abstract\n\n(.+?)\n", md, re.S).group(1).strip()
abs_tbl = [t for t in hdr._element.iter(W + "tbl") if "Summarize your project" in "".join(x.text or "" for x in t.iter(W + "t"))][-1]
abs_p = docx.text.paragraph.Paragraph(abs_tbl.find(".//" + W + "p"), hdr._parent)
for r in abs_p.runs:
    r._element.getparent().remove(r._element)
add_inline(abs_p, abstract)
abs_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Drop the guidance box and all guidance text after the header block.
sectpr = body.find(W + "sectPr")
for el in list(body):
    if el is not hdr._element and el is not sectpr:
        body.remove(el)

# Body: everything after the abstract.
rest = md.split("## Abstract", 1)[1].split("\n## ", 1)[1]
rest = "## " + rest
blocks = [b for b in re.split(r"\n\s*\n", rest) if b.strip()]
heading_tpl = None


def heading(text, level):
    p = d.add_paragraph()
    p.style = d.styles["Heading 2" if level == 2 else "Heading 3"]
    spacing(p, before=12 if level == 2 else 8, after=4, line=1.0)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), 14 if level == 2 else 12, bold=True)
    p.runs[0].font.color.rgb = docx.shared.RGBColor(0, 0, 0)


def border(cell, edges):
    tcpr = cell._element.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single" if edge in edges else "nil")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "999999")
        b.append(e)
    tcpr.append(b)


def table(lines):
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines if not re.match(r"^\|[-| ]+\|$", ln.strip())]
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [1.55] + [(6.5 - 1.55) / (len(rows[0]) - 1)] * (len(rows[0]) - 1)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Inches(widths[j])
            p = cell.paragraphs[0]
            spacing(p, 0, 0, 1.0)
            set_font(p.add_run(val), 7.5, bold=(i == 0) or None)
            border(cell, {"bottom"} if i in (0, len(rows) - 1) else set())
    after = d.add_paragraph()
    spacing(after, 0, 4, 1.0)


for b in blocks:
    lines = b.strip().split("\n")
    first = lines[0]
    if first.startswith("### "):
        heading(first[4:], 3)
    elif first.startswith("## "):
        heading(first[3:], 2)
    elif first.startswith("!["):
        src = re.search(r"\((.+?)\)", first).group(1)
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(p, 4, 2, 1.0)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ROOT / "report" / src), width=Inches(5.9))
    elif first.startswith("|"):
        table(lines)
    elif re.match(r"^(- |\d+\. )", first):
        items, cur = [], None
        for ln in lines:
            if re.match(r"^(- |\d+\. )", ln):
                cur = re.sub(r"^(- |\d+\. )", "", ln)
                items.append(cur)
            else:
                items[-1] += " " + ln.strip()
        is_ref = re.match(r"^\d+\. [A-Z][A-Za-z-]+, ", first) is not None
        num = NUM_BULLET if first.startswith("- ") else (NUM_REFS if is_ref else NUM_CONTRIB)
        for it in items:
            p = d.add_paragraph()
            numbered(p, num)
            spacing(p, 0, 3 if not is_ref else 2, 1.15)
            add_inline(p, it, size=10 if is_ref else None)
    else:
        text = " ".join(ln.strip() for ln in lines)
        cap = re.match(r"^\*(Figure|Table) [A-Z]?\d", text) is not None
        p = d.add_paragraph()
        spacing(p, 0, 8, 1.15 if not cap else 1.0)
        if cap and text.startswith("*Table"):
            p.paragraph_format.keep_with_next = True
        add_inline(p, text.strip("*") if cap else text, size=9.5 if cap else None, italic=cap)

body.remove(sectpr)
body.append(sectpr)
out = ROOT / "report/report.docx"
d.save(out)
print("wrote", out)
